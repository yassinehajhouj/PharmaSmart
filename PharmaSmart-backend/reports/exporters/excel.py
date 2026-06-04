"""
Excel export using openpyxl — produces a styled .xlsx file from report data.
"""
from __future__ import annotations

from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter


# ── style constants ───────────────────────────────────────────────────────────

_GREEN    = '0B6E4F'
_DARK     = '1B4965'
_ACCENT   = '14A76C'
_LIGHT_BG = 'F0F7F4'
_HEADER   = PatternFill('solid', fgColor=_GREEN)
_SUBHEAD  = PatternFill('solid', fgColor=_DARK)
_ALT_ROW  = PatternFill('solid', fgColor='E8F5F0')
_WHITE_FT = Font(color='FFFFFF', bold=True, size=11)
_BOLD     = Font(bold=True)
_THIN     = Side(border_style='thin', color='CCCCCC')
_BORDER   = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

MAD_FMT   = '#,##0.00" MAD"'
INT_FMT   = '#,##0'


def _h(ws, row, col, value, fill=None, font=None, align='left', num_fmt=None):
    cell = ws.cell(row=row, column=col, value=value)
    if fill:   cell.fill      = fill
    if font:   cell.font      = font
    cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    cell.border = _BORDER
    if num_fmt: cell.number_format = num_fmt
    return cell


def _autowidth(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or '')))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)


# ── main builder ──────────────────────────────────────────────────────────────

def generate_excel(pharmacie, period: str, report_type: str, data: dict) -> BytesIO:
    wb = Workbook()

    _build_summary_sheet(wb.active, pharmacie, period, report_type, data)

    if report_type in ('general', 'sales'):
        _build_sales_sheet(wb.create_sheet('Ventes détaillées'), data)

    if report_type in ('general', 'stock'):
        _build_stock_sheet(wb.create_sheet('Stock'), data)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── sheet builders ────────────────────────────────────────────────────────────

def _build_summary_sheet(ws, pharmacie, period, report_type, data):
    ws.title = 'Résumé'
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 40

    PERIOD_LABELS = {'month': 'Ce mois', 'quarter': 'Ce trimestre', 'year': 'Cette année'}

    # ── Title banner ──
    ws.merge_cells('A1:G1')
    title_cell = ws['A1']
    title_cell.value = f'Rapport {report_type.upper()} — {PERIOD_LABELS.get(period, period)}'
    title_cell.fill  = _HEADER
    title_cell.font  = Font(color='FFFFFF', bold=True, size=14)
    title_cell.alignment = Alignment(horizontal='center', vertical='center')

    # ── Meta row ──
    ws.merge_cells('A2:G2')
    meta = ws['A2']
    meta.value = (
        f'Pharmacie: {pharmacie.nom_pharmacie}  |  '
        f'Généré le: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    )
    meta.font = Font(italic=True, color='555555')
    meta.alignment = Alignment(horizontal='center')

    row = 4

    # ── KPIs ──
    kpis = data.get('kpis', {})
    _h(ws, row, 1, 'Indicateurs clés', fill=_SUBHEAD, font=_WHITE_FT)
    ws.merge_cells(f'A{row}:G{row}')
    row += 1

    kpi_items = [
        ('Chiffre d\'affaires',   kpis.get('ca', 0),                    MAD_FMT),
        ('Commandes traitées',    kpis.get('commandes', 0),              INT_FMT),
        ('Médicaments vendus',    kpis.get('medicaments_vendus', 0),     INT_FMT),
        ('Valeur du stock',       kpis.get('valeur_stock', 0),           MAD_FMT),
    ]
    for label, value, fmt in kpi_items:
        _h(ws, row, 1, label, font=_BOLD)
        _h(ws, row, 2, value, num_fmt=fmt, align='right')
        ws.merge_cells(f'B{row}:G{row}')
        row += 1

    row += 1

    # ── Revenue evolution ──
    rev = data.get('revenue_evolution', [])
    if rev:
        _h(ws, row, 1, 'Evolution du CA (6 derniers mois)', fill=_SUBHEAD, font=_WHITE_FT)
        ws.merge_cells(f'A{row}:G{row}')
        row += 1
        _h(ws, row, 1, 'Mois',  fill=_HEADER, font=_WHITE_FT)
        _h(ws, row, 2, 'CA (MAD)', fill=_HEADER, font=_WHITE_FT, align='right')
        row += 1
        for i, m in enumerate(rev):
            fill = _ALT_ROW if i % 2 == 0 else None
            _h(ws, row, 1, m.get('mois', ''), fill=fill)
            _h(ws, row, 2, m.get('ca', 0), fill=fill, num_fmt=MAD_FMT, align='right')
            row += 1
        row += 1

    # ── Top medicines ──
    top = data.get('top_medicaments', [])
    if top:
        _h(ws, row, 1, 'Top médicaments vendus', fill=_SUBHEAD, font=_WHITE_FT)
        ws.merge_cells(f'A{row}:G{row}')
        row += 1
        headers = ['#', 'Médicament', 'Catégorie', 'Qté vendue', 'Revenu (MAD)']
        for j, h in enumerate(headers, 1):
            _h(ws, row, j, h, fill=_HEADER, font=_WHITE_FT, align='center')
        row += 1
        for i, med in enumerate(top, 1):
            fill = _ALT_ROW if i % 2 == 0 else None
            _h(ws, row, 1, i, fill=fill, align='center')
            _h(ws, row, 2, med.get('medicament__nom', ''), fill=fill)
            _h(ws, row, 3, med.get('medicament__categorie__nom', '—'), fill=fill)
            _h(ws, row, 4, med.get('total_vendu', 0), fill=fill, num_fmt=INT_FMT, align='right')
            _h(ws, row, 5, float(med.get('total_revenu', 0)), fill=fill, num_fmt=MAD_FMT, align='right')
            row += 1

    _autowidth(ws)


def _build_sales_sheet(ws, data):
    ws.sheet_view.showGridLines = False
    _h(ws, 1, 1, 'Médicament', fill=_HEADER, font=_WHITE_FT)
    _h(ws, 1, 2, 'Catégorie',  fill=_HEADER, font=_WHITE_FT)
    _h(ws, 1, 3, 'Qté vendue', fill=_HEADER, font=_WHITE_FT, align='right')
    _h(ws, 1, 4, 'Revenu',     fill=_HEADER, font=_WHITE_FT, align='right')

    for i, med in enumerate(data.get('top_medicaments', []), 2):
        fill = _ALT_ROW if i % 2 == 0 else None
        _h(ws, i, 1, med.get('medicament__nom', ''), fill=fill)
        _h(ws, i, 2, med.get('medicament__categorie__nom', '—'), fill=fill)
        _h(ws, i, 3, med.get('total_vendu', 0), fill=fill, num_fmt=INT_FMT, align='right')
        _h(ws, i, 4, float(med.get('total_revenu', 0)), fill=fill, num_fmt=MAD_FMT, align='right')
    _autowidth(ws)


def _build_stock_sheet(ws, data):
    ws.sheet_view.showGridLines = False
    stock_items = data.get('stock_details', [])
    headers = ['Médicament', 'Catégorie', 'Stock', 'Seuil alerte', 'Prix vente', 'Valeur (MAD)', 'Statut']
    for j, h in enumerate(headers, 1):
        _h(ws, 1, j, h, fill=_HEADER, font=_WHITE_FT)
    for i, s in enumerate(stock_items, 2):
        fill = _ALT_ROW if i % 2 == 0 else None
        _h(ws, i, 1, s.get('nom', ''), fill=fill)
        _h(ws, i, 2, s.get('categorie', '—'), fill=fill)
        _h(ws, i, 3, s.get('quantite', 0), fill=fill, num_fmt=INT_FMT, align='right')
        _h(ws, i, 4, s.get('seuil_alerte', 0), fill=fill, num_fmt=INT_FMT, align='right')
        _h(ws, i, 5, float(s.get('prix_vente', 0)), fill=fill, num_fmt=MAD_FMT, align='right')
        _h(ws, i, 6, float(s.get('valeur', 0)), fill=fill, num_fmt=MAD_FMT, align='right')
        _h(ws, i, 7, s.get('statut', ''), fill=fill)
    _autowidth(ws)
